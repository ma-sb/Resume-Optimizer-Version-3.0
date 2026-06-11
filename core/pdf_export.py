"""Reconstruct exported resume styling for PDF-origin resumes."""

from __future__ import annotations

import json
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from reportlab.lib import colors
from reportlab.lib.pagesizes import LETTER
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

from .matcher import _make_client
from .reader import Resume

_PDF_STYLE_SYSTEM = """You classify the visual style of a resume that originated from PDF.
Choose the closest reconstruction profile for export.

Return ONLY valid JSON:
{
  "template": "classic" | "modern" | "compact",
  "name_alignment": "left" | "center",
  "heading_case": "upper" | "title",
  "heading_rule": true | false,
  "density": "compact" | "standard" | "spacious",
  "accent": "navy" | "slate" | "teal" | "none"
}

Use these rules:
- classic: conservative, standard spacing, minimal decoration
- modern: cleaner spacing, stronger accent, centered or visually airy
- compact: dense spacing, tighter margins, more ATS-like, many short lines
- If unsure, choose classic
- Do not invent any other values
"""

_DEFAULT_STYLE = {
    "template": "classic",
    "name_alignment": "left",
    "heading_case": "upper",
    "heading_rule": True,
    "density": "standard",
    "accent": "navy",
}

_ACCENTS = {
    "navy": {"rgb": RGBColor(15, 52, 96), "rl": colors.HexColor("#0F3460"), "hex": "#0F3460"},
    "slate": {"rgb": RGBColor(51, 65, 85), "rl": colors.HexColor("#334155"), "hex": "#334155"},
    "teal": {"rgb": RGBColor(13, 148, 136), "rl": colors.HexColor("#0D9488"), "hex": "#0D9488"},
    "none": {"rgb": RGBColor(17, 24, 39), "rl": colors.HexColor("#111827"), "hex": "#111827"},
}


def _parse_json(text: str) -> dict:
    cleaned = re.sub(r"^```(?:json)?\s*\n?", "", text.strip())
    cleaned = re.sub(r"\n?```\s*$", "", cleaned)
    cleaned = re.sub(r"[\x00-\x09\x0b\x0c\x0e-\x1f]", " ", cleaned)
    start = cleaned.index("{")
    obj, _ = json.JSONDecoder().raw_decode(cleaned[start:])
    return obj


def _call_style_classifier(
    api_key: str,
    model: str,
    base_url: str,
    json_mode: bool,
    prompt: str,
) -> dict:
    client = _make_client(api_key, base_url)
    kwargs: dict = {
        "model": model,
        "temperature": 0.1,
        "messages": [
            {"role": "system", "content": _PDF_STYLE_SYSTEM},
            {"role": "user", "content": prompt},
        ],
    }
    if json_mode:
        kwargs["response_format"] = {"type": "json_object"}
    response = client.chat.completions.create(**kwargs)
    return _parse_json(response.choices[0].message.content or "")


def infer_pdf_style_profile(
    resume: Resume,
    api_key: str,
    model: str,
    base_url: str = "",
    json_mode: bool = True,
) -> dict:
    """Infer the closest reconstruction style for a PDF resume."""
    lines = [line.strip() for line in resume.full_text.splitlines() if line.strip()]
    short_lines = sum(1 for line in lines if len(line) <= 32)
    avg_len = round(sum(len(line) for line in lines) / max(1, len(lines)), 1)
    heading_like = [line for line in lines if len(line) <= 28 and (line.isupper() or line.istitle())][:12]
    sample = "\n".join(lines[:80])

    prompt = f"""
FILENAME: {resume.filename}
CANDIDATE: {resume.name or "Unknown"}

LAYOUT HINTS:
- total_nonempty_lines: {len(lines)}
- short_lines_under_33_chars: {short_lines}
- average_line_length: {avg_len}
- heading_like_lines: {heading_like}

RESUME TEXT SAMPLE:
{sample}
""".strip()

    try:
        raw = _call_style_classifier(api_key, model, base_url, json_mode, prompt)
    except Exception:
        raw = {}

    style = dict(_DEFAULT_STYLE)
    style.update({k: v for k, v in raw.items() if k in style})
    if style["accent"] not in _ACCENTS:
        style["accent"] = _DEFAULT_STYLE["accent"]
    if style["template"] not in {"classic", "modern", "compact"}:
        style["template"] = _DEFAULT_STYLE["template"]
    if style["name_alignment"] not in {"left", "center"}:
        style["name_alignment"] = _DEFAULT_STYLE["name_alignment"]
    if style["heading_case"] not in {"upper", "title"}:
        style["heading_case"] = _DEFAULT_STYLE["heading_case"]
    if style["density"] not in {"compact", "standard", "spacious"}:
        style["density"] = _DEFAULT_STYLE["density"]
    style["heading_rule"] = bool(style["heading_rule"])
    return style


def _safe_filename(text: str) -> str:
    return (text or "").strip().replace(" ", "_").replace("/", "_")


def _file_stem(name: str, company_name: str) -> str:
    parts = name.split()
    first, last = (parts[0], parts[-1]) if len(parts) >= 2 else (name, "")
    return f"{_safe_filename(first)}_{_safe_filename(last)}_Resume_{_safe_filename(company_name or 'Company')}"


def _heading_text(heading: str, case: str) -> str:
    return heading.upper() if case == "upper" else heading.title()


def _doc_spacing(density: str) -> tuple[float, float]:
    return {
        "compact": (0.05, 0.02),
        "spacious": (0.14, 0.08),
        "standard": (0.1, 0.04),
    }.get(density, (0.1, 0.04))


def export_pdf_reconstructed_docx(
    *,
    optimized: dict,
    output_dir: str | Path,
    company_name: str,
    style_profile: dict,
) -> Path:
    """Generate a reconstructed .docx when the source resume was a PDF."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    accent_rgb = _ACCENTS[style_profile["accent"]]["rgb"]
    doc = Document()
    section = doc.sections[0]
    if style_profile["template"] == "compact":
        margin = Inches(0.55)
    elif style_profile["template"] == "modern":
        margin = Inches(0.75)
    else:
        margin = Inches(0.7)
    section.top_margin = margin
    section.bottom_margin = margin
    section.left_margin = margin
    section.right_margin = margin

    name_p = doc.add_paragraph()
    name_run = name_p.add_run((optimized.get("name") or "Candidate").strip())
    name_run.bold = True
    name_run.font.size = Pt(18 if style_profile["template"] != "compact" else 16)
    name_run.font.color.rgb = accent_rgb
    name_p.alignment = WD_ALIGN_PARAGRAPH.CENTER if style_profile["name_alignment"] == "center" else WD_ALIGN_PARAGRAPH.LEFT
    before, after = _doc_spacing(style_profile["density"])
    name_p.paragraph_format.space_after = Pt(10 if style_profile["density"] != "compact" else 6)

    for sec in optimized.get("sections", []):
        heading = (sec.get("heading") or "").strip()
        if not heading:
            continue

        heading_p = doc.add_paragraph()
        heading_run = heading_p.add_run(_heading_text(heading, style_profile["heading_case"]))
        heading_run.bold = True
        heading_run.font.size = Pt(11 if style_profile["template"] == "compact" else 12)
        heading_run.font.color.rgb = accent_rgb
        heading_p.paragraph_format.space_before = Pt(10 if style_profile["density"] != "compact" else 7)
        heading_p.paragraph_format.space_after = Pt(4 if style_profile["density"] != "compact" else 2)
        if style_profile["heading_rule"]:
            border_p = doc.add_paragraph("_" * 36)
            border_p.runs[0].font.color.rgb = accent_rgb
            border_p.paragraph_format.space_before = Pt(0)
            border_p.paragraph_format.space_after = Pt(4 if style_profile["density"] != "compact" else 2)

        content = (sec.get("content") or "").strip()
        if content:
            p = doc.add_paragraph(content)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4 if style_profile["density"] != "compact" else 2)

        for bullet in sec.get("bullets", []):
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(bullet)
            p.paragraph_format.space_before = Pt(before * 72)
            p.paragraph_format.space_after = Pt(after * 72)

    stem = _file_stem((optimized.get("name") or "Candidate").strip(), company_name)
    path = output_dir / f"{stem}.docx"
    doc.save(str(path))
    return path


def export_pdf_reconstructed_pdf(
    *,
    optimized: dict,
    output_dir: str | Path,
    company_name: str,
    style_profile: dict,
) -> Path:
    """Generate a reconstructed PDF when the source resume was a PDF."""
    output_dir = Path(output_dir)
    output_dir.mkdir(parents=True, exist_ok=True)

    stem = _file_stem((optimized.get("name") or "Candidate").strip(), company_name)
    path = output_dir / f"{stem}.pdf"
    accent = _ACCENTS[style_profile["accent"]]["rl"]
    accent_hex = _ACCENTS[style_profile["accent"]]["hex"]

    margin = {
        "compact": 0.5 * inch,
        "spacious": 0.85 * inch,
        "standard": 0.7 * inch,
    }.get(style_profile["density"], 0.7 * inch)
    doc = SimpleDocTemplate(
        str(path),
        pagesize=LETTER,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
    )

    sample = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "ResumeName",
        parent=sample["Title"],
        fontName="Helvetica-Bold",
        fontSize=18 if style_profile["template"] != "compact" else 16,
        textColor=accent,
        alignment=1 if style_profile["name_alignment"] == "center" else 0,
        spaceAfter=12 if style_profile["density"] != "compact" else 8,
    )
    heading_style = ParagraphStyle(
        "ResumeHeading",
        parent=sample["Heading2"],
        fontName="Helvetica-Bold",
        fontSize=11 if style_profile["template"] == "compact" else 12,
        textColor=accent,
        spaceBefore=10 if style_profile["density"] != "compact" else 7,
        spaceAfter=4 if style_profile["density"] != "compact" else 2,
    )
    body_style = ParagraphStyle(
        "ResumeBody",
        parent=sample["BodyText"],
        fontName="Helvetica",
        fontSize=9.5 if style_profile["template"] == "compact" else 10.2,
        leading=12 if style_profile["density"] == "compact" else 13.5,
        textColor=colors.black,
        spaceAfter=4 if style_profile["density"] != "compact" else 2,
    )
    bullet_style = ParagraphStyle(
        "ResumeBullet",
        parent=body_style,
        leftIndent=12,
        firstLineIndent=-8,
    )

    story = [Paragraph(escape((optimized.get("name") or "Candidate").strip()), name_style)]
    for sec in optimized.get("sections", []):
        heading = (sec.get("heading") or "").strip()
        if not heading:
            continue
        story.append(Paragraph(escape(_heading_text(heading, style_profile["heading_case"])), heading_style))
        if style_profile["heading_rule"]:
            story.append(Paragraph(f'<font color="{accent_hex}">{"_" * 40}</font>', body_style))
        content = (sec.get("content") or "").strip()
        if content:
            story.append(Paragraph(escape(content), body_style))
        for bullet in sec.get("bullets", []):
            story.append(Paragraph("&bull; " + escape(bullet), bullet_style))
        story.append(Spacer(1, 3 if style_profile["density"] == "compact" else 5))

    doc.build(story)
    return path


def export_pdf_origin_resume(
    *,
    optimized: dict,
    source_resume: Resume,
    output_dir: str | Path,
    company_name: str,
    api_key: str,
    model: str,
    base_url: str = "",
    json_mode: bool = True,
) -> tuple[Path, Path, dict]:
    """Reconstruct Word and PDF exports for a PDF-origin resume."""
    style_profile = infer_pdf_style_profile(
        source_resume,
        api_key=api_key,
        model=model,
        base_url=base_url,
        json_mode=json_mode,
    )
    docx_path = export_pdf_reconstructed_docx(
        optimized=optimized,
        output_dir=output_dir,
        company_name=company_name,
        style_profile=style_profile,
    )
    pdf_path = export_pdf_reconstructed_pdf(
        optimized=optimized,
        output_dir=output_dir,
        company_name=company_name,
        style_profile=style_profile,
    )
    return docx_path, pdf_path, style_profile
